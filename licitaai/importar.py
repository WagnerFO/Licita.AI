"""Importa os contratos da planilha para o banco (serve como dados de teste)."""
import re
from datetime import datetime

import click
from flask.cli import with_appcontext
from openpyxl import load_workbook

from .db import get_db


def texto(valor):
    if valor is None:
        return ""
    valor = str(valor).strip()
    return "" if valor == "-" else valor


def para_centavos(valor):
    if isinstance(valor, (int, float)):
        return round(valor * 100)
    return None


def para_data(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return None


@click.command("importar-planilha")
@click.argument("arquivo", type=click.Path(exists=True))
@with_appcontext
def importar_command(arquivo):
    """Importa a aba 'Relações dos Contratos' de um arquivo .xlsx/.xlsm."""
    planilha = load_workbook(arquivo, read_only=True, data_only=True)
    aba = planilha["Relações dos Contratos"]
    db = get_db()
    importados = 0
    ignorados = 0

    # os dados começam na linha 5; limitamos linhas e colunas por segurança
    for linha in aba.iter_rows(min_row=5, max_row=2000, max_col=25, values_only=True):
        numero = texto(linha[10])
        if not numero:
            continue

        data_assinatura = para_data(linha[0])
        data_fim = para_data(linha[1])
        # em ATAs o valor fica na coluna "Valor Total da ATA"
        valor_total = para_centavos(linha[20])
        if valor_total is None:
            valor_total = para_centavos(linha[21])

        if not data_assinatura or not data_fim or valor_total is None:
            ignorados += 1
            continue

        prazo = re.search(r"\d+", texto(linha[2]))
        nao_informado = "Não informado"

        try:
            db.execute(
                """INSERT INTO contratos
                   (numero, processo_licitatorio, numero_modalidade, modalidade, empresa,
                    cnpj, resumo_objeto, objeto, valor_estimado_centavos,
                    valor_total_centavos, orgao_gestor, secretario, data_assinatura,
                    data_fim, prazo_vigencia_meses)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (numero, texto(linha[11]) or nao_informado, texto(linha[12]),
                 texto(linha[13]) or nao_informado, texto(linha[14]) or nao_informado,
                 texto(linha[15]) or nao_informado, texto(linha[16]) or nao_informado,
                 texto(linha[17]) or nao_informado, para_centavos(linha[19]),
                 valor_total, texto(linha[23]) or nao_informado, texto(linha[24]),
                 data_assinatura, data_fim, int(prazo.group()) if prazo else None),
            )
            importados += 1
        except Exception:
            ignorados += 1  # número de contrato repetido ou dado inválido

    db.commit()
    click.echo(f"Importados: {importados} | Ignorados: {ignorados}")
